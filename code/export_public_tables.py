#!/usr/bin/env python3
"""Export the public machine-readable tables from the frozen workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "outputs" / "TRUSTWORTHY_ORCHESTRATION_REVIEW_EVIDENCE_WORKBOOK.xlsx"
OUT = ROOT / "data"
EXPORTS = {
    "07_TA_SCREENING": "title_abstract_screening.csv",
    "08_FULL_TEXT_SCREENING": "full_text_screening.csv",
    "09_INCLUDED_STUDIES": "included_studies.csv",
    "11_EVIDENCE_MAP": "evidence_map.csv",
    "12_QUALITY_APPRAISAL": "quality_appraisal.csv",
}


def cleaned_rows(ws):
    rows = []
    for number, row in enumerate(ws.iter_rows(values_only=True), 1):
        if number < 4:
            continue
        values = ["" if value is None else value for value in row]
        while values and values[-1] == "":
            values.pop()
        if values:
            rows.append(values)
    return rows


def main() -> None:
    OUT.mkdir(exist_ok=True)
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    for sheet_name, filename in EXPORTS.items():
        rows = cleaned_rows(workbook[sheet_name])
        width = len(rows[0])
        destination = OUT / filename
        with destination.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerows([row + [""] * (width - len(row)) for row in rows])
        print(f"{sheet_name}: {len(rows) - 1} records -> {destination.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
