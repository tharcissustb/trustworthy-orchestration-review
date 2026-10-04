#!/usr/bin/env python3
"""Validate the public workbook's structure, counts, keys, and frozen results."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "outputs" / "TRUSTWORTHY_ORCHESTRATION_REVIEW_EVIDENCE_WORKBOOK.xlsx"
EXPECTED_SHEETS = [
    "00_READ_ME", "01_DASHBOARD", "02_PRISMA_COUNTS", "03_SEARCH_STRATEGY",
    "04_RAW_SEARCH_DATA", "05_DEDUPLICATED_RECORDS", "06_SCREENING_CRITERIA",
    "07_TA_SCREENING", "08_FULL_TEXT_SCREENING", "09_INCLUDED_STUDIES",
    "10_CODEBOOK", "11_EVIDENCE_MAP", "12_QUALITY_APPRAISAL",
    "13_RELIABILITY", "14_VALIDATION_CASE", "15_DATA_DICTIONARY",
]
ERROR_TOKENS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def table(ws, header_row=4):
    rows = []
    for number, row in enumerate(ws.iter_rows(values_only=True), 1):
        if number < header_row:
            continue
        values = list(row)
        while values and values[-1] is None:
            values.pop()
        if values:
            rows.append(values)
    return rows


def fail_if(condition, message, errors):
    if condition:
        errors.append(message)


def main() -> None:
    errors = []
    with ZipFile(PATH) as archive:
        fail_if(archive.testzip() is not None, "XLSX ZIP integrity check failed", errors)

    workbook = load_workbook(PATH, read_only=True, data_only=True)
    fail_if(workbook.sheetnames != EXPECTED_SHEETS, "Unexpected sheet order", errors)

    raw = table(workbook["04_RAW_SEARCH_DATA"])
    dedup = table(workbook["05_DEDUPLICATED_RECORDS"])
    ta = table(workbook["07_TA_SCREENING"])
    full_text = table(workbook["08_FULL_TEXT_SCREENING"])
    included = table(workbook["09_INCLUDED_STUDIES"])
    evidence = table(workbook["11_EVIDENCE_MAP"])
    quality = table(workbook["12_QUALITY_APPRAISAL"])

    expected_counts = {
        "raw": (len(raw) - 1, 20930),
        "deduplicated": (len(dedup) - 1, 14597),
        "title_abstract": (len(ta) - 1, 757),
        "full_text": (len(full_text) - 1, 529),
        "included": (len(included) - 1, 120),
        "evidence": (len(evidence) - 1, 120),
        "quality": (len(quality) - 1, 120),
    }
    for name, (actual, expected) in expected_counts.items():
        fail_if(actual != expected, f"{name}: expected {expected}, found {actual}", errors)

    inc_keys = [str(row[1]) for row in included[1:]]
    ev_keys = [str(row[0]) for row in evidence[1:]]
    qa_keys = [str(row[0]) for row in quality[1:]]
    fail_if(len(set(inc_keys)) != 120, "Included-study keys are not unique", errors)
    fail_if(set(inc_keys) != set(ev_keys) or set(inc_keys) != set(qa_keys), "Included/evidence/quality key sets differ", errors)

    years = Counter(str(row[4]) for row in included[1:])
    expected_years = {"2021": 10, "2022": 11, "2023": 19, "2024": 28, "2025": 44, "2026": 8}
    fail_if(dict(sorted(years.items())) != expected_years, f"Unexpected year distribution: {dict(years)}", errors)

    full_text_headers = {name: index for index, name in enumerate(full_text[0])}
    decisions = Counter(str(row[full_text_headers["FT_decision"]]) for row in full_text[1:])
    expected_decisions = {"Excluded after full text": 488, "Include in analytical extension": 36, "Not retrieved": 5}
    fail_if(dict(decisions) != expected_decisions, f"Unexpected full-text decisions: {dict(decisions)}", errors)

    spreadsheet_errors = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str) and value in ERROR_TOKENS:
                    spreadsheet_errors.append(f"{ws.title}:{value}")
    fail_if(bool(spreadsheet_errors), f"Spreadsheet errors: {spreadsheet_errors[:10]}", errors)

    result = {
        "status": "PASS" if not errors else "FAIL",
        "errors": errors,
        "counts": {name: actual for name, (actual, _) in expected_counts.items()},
        "year_distribution": expected_years,
        "full_text_decisions": expected_decisions,
        "key_sets_equal": set(inc_keys) == set(ev_keys) == set(qa_keys),
    }
    print(json.dumps(result, indent=2))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
