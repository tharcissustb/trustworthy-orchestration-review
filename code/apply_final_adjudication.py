#!/usr/bin/env python3
"""Apply the frozen 16-study P1–P8 consensus and cross-route reconciliation."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
ADJUDICATION = ROOT / "data" / "p1_p8_reliability_adjudication.xlsx"
REDISCOVERIES = {5, 1541, 13877, 2852, 10858, 3263, 9957, 9948, 3277, 4899}


def read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


def write_csv(path: Path, fields: list[str], rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def final_codes() -> dict[tuple[str, str], int]:
    workbook = load_workbook(ADJUDICATION, read_only=True, data_only=True)
    sheet = workbook["02_FINAL_SAMPLE_CODES"]
    rows = sheet.iter_rows(min_row=4, values_only=True)
    header = next(rows)
    index = {name: position for position, name in enumerate(header)}
    result: dict[tuple[str, str], int] = {}
    for row in rows:
        if not row[index["Sample ID"]]:
            continue
        key = str(row[index["Bibliographic key"]])
        mechanism = str(row[index["Mechanism"]])
        result[(key, mechanism)] = int(row[index["Final code"]])
    if len(result) != 128:
        raise RuntimeError(f"Expected 128 final mechanism codes; found {len(result)}")
    return result


def main() -> None:
    codes = final_codes()

    evidence_path = ROOT / "data" / "evidence_map.csv"
    fields, rows = read_csv(evidence_path)
    changes = 0
    for row in rows:
        for number in range(1, 9):
            mechanism = f"P{number}"
            key = (row["bib_key"], mechanism)
            if key not in codes:
                continue
            value = str(codes[key])
            if row[mechanism] != value:
                row[mechanism] = value
                changes += 1
    if changes != 51:
        raise RuntimeError(f"Expected 51 evidence-map changes; found {changes}")
    write_csv(evidence_path, fields, rows)

    full_text_path = ROOT / "data" / "full_text_screening.csv"
    fields, rows = read_csv(full_text_path)
    changed = 0
    for row in rows:
        if int(row["record_id"]) not in REDISCOVERIES:
            continue
        row["FT_decision"] = "Eligible — cross-route rediscovery"
        row["FT_exclusion_code"] = "FT-RD"
        row["decision_evidence_location"] = row["decision_evidence_location"] or "Existing included-study register"
        row["notes"] = "Eligible after full-text assessment; already present in the citation-led set and counted once during cross-route reconciliation."
        changed += 1
    if changed != 10:
        raise RuntimeError(f"Expected 10 full-text rediscoveries; found {changed}")
    write_csv(full_text_path, fields, rows)

    ta_path = ROOT / "data" / "title_abstract_screening.csv"
    fields, rows = read_csv(ta_path)
    changed = 0
    for row in rows:
        if int(row["record_id"]) in REDISCOVERIES:
            row["corpus_disposition"] = "Eligible — cross-route rediscovery"
            changed += 1
    if changed != 10:
        raise RuntimeError(f"Expected 10 title/abstract rediscoveries; found {changed}")
    write_csv(ta_path, fields, rows)

    print("Applied 51 adjudicated mechanism changes and 10 cross-route rediscoveries.")


if __name__ == "__main__":
    main()
